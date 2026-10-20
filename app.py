import streamlit as st
import pandas as pd
import numpy as np
import io
import re
import textwrap
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="DS Transfer Summary", layout="wide")

# Custom CSS Theme
st.markdown("""
    <style>
    .stApp {
        background-color: #0E1117;
        color: #FFFFFF;
    }
    
    .custom-table {
        width: 100%;
        border-collapse: collapse;
        font-family: Arial, sans-serif;
        font-size: 14px;
        color: #000000;
        margin-top: 15px;
        background-color: #FFFFFF;
    }
    .custom-table th {
        background-color: #EAEAEA !important;
        color: #000000 !important;
        font-weight: bold;
        border: 1px solid #D1D5DB;
        text-align: center;
        padding: 10px;
    }
    .custom-table td {
        border: 1px solid #E5E7EB;
        padding: 8px 12px;
        color: #000000;
    }
    .custom-table tr.subtotal-row td {
        background-color: #F3F4F6 !important;
        color: #000000 !important;
        font-weight: bold !important;
        border-top: 1.5px solid #9CA3AF;
        border-bottom: 1.5px solid #9CA3AF;
    }
    .custom-table tr.total-row td {
        background-color: #F8B486 !important;
        color: #000000 !important;
        font-weight: bold !important;
        border-top: 2px solid #D97706;
    }
    
    div[data-testid="stMetricValue"] {
        font-size: 30px;
        color: #F8B486;
        font-weight: bold;
    }
    div[data-testid="stMetricLabel"] {
        color: #E5E7EB;
        font-weight: 600;
    }
    section[data-testid="stSidebar"] {
        background-color: #1F2937;
        color: #FFFFFF;
    }
    </style>
""", unsafe_allow_html=True)

st.title("📦 Dispatched Quantity Summary (DS Wise)")

uploaded_file = st.sidebar.file_uploader("Upload Transfer Order File", type=["csv", "xlsx"])
use_demo = st.sidebar.checkbox("Use Demo Data", value=(uploaded_file is None))

def transform_reason(reason_str):
    reason_str = str(reason_str).strip()
    
    # 1. Match 'HomeRun picker · TO-XXXXX · vehicle YYYYY'
    picker_match = re.search(r'HomeRun picker\s*·\s*(TO-\d+)', reason_str, re.IGNORECASE)
    if picker_match:
        to_num = picker_match.group(1).upper()
        return f"picked by picker - {to_num}"
    
    # 2. Check for cement
    if re.search(r'cement', reason_str, re.IGNORECASE):
        return "BAG"
    
    # 3. Check for ply / internal transfer
    if re.search(r'ply|plywood|internal transfer', reason_str, re.IGNORECASE):
        return "PLY"
        
    return reason_str

def generate_styled_excel(rows_list, grand_total):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sent_Summary"

    headers = ["date", "DS", "TO", "reson", "qty"]
    ws.append(headers)

    for row in rows_list:
        ws.append([row['date'], row['DS'], row['TO'], row['reson'], row['qty']])

    # Grand Total Row
    ws.append(["Total", "", "", "Total Dispatched Qty", grand_total])

    font_header = Font(name="Arial", size=11, bold=True, color="000000")
    font_body = Font(name="Arial", size=10, color="000000")
    font_subtotal = Font(name="Arial", size=10, bold=True, color="000000")
    font_total = Font(name="Arial", size=11, bold=True, color="000000")

    fill_header = PatternFill(start_color="EAEAEA", end_color="EAEAEA", fill_type="solid")
    fill_even = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    fill_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_subtotal = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
    fill_total = PatternFill(start_color="F8B486", end_color="F8B486", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    # Format Headers
    for col_num in range(1, 6):
        cell = ws.cell(row=1, column=col_num)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    max_row = ws.max_row
    for row_num in range(2, max_row):
        is_subtotal = ws.cell(row=row_num, column=4).value == "Sub-Total Qty"
        fill = fill_subtotal if is_subtotal else (fill_even if row_num % 2 == 0 else fill_odd)
        font = font_subtotal if is_subtotal else font_body

        for col_num in range(1, 6):
            cell = ws.cell(row=row_num, column=col_num)
            cell.font = font
            cell.fill = fill
            cell.border = thin_border
            if col_num in [1, 3]:
                cell.alignment = Alignment(horizontal="center")
            elif col_num == 5:
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = '#,##0'
            else:
                cell.alignment = Alignment(horizontal="left")

    # Format Grand Total Row
    for col_num in range(1, 6):
        cell = ws.cell(row=max_row, column=col_num)
        cell.font = font_total
        cell.fill = fill_total
        cell.border = thin_border
        if col_num == 5:
            cell.alignment = Alignment(horizontal="right")
            cell.number_format = '#,##0'
        elif col_num in [1, 4]:
            cell.alignment = Alignment(horizontal="left")

    # Column Widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()

df = None

if uploaded_file is not None:
    df = pd.read_csv(uploaded_file) if uploaded_file.name.endswith('.csv') else pd.read_excel(uploaded_file)
elif use_demo:
    df = pd.DataFrame({
        "Date": ["2026-10-10"] * 7,
        "From Location Name": ["DC01 Rampura"] * 7
